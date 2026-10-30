# Text for character drafts. Plain text and Markdown are read as they are. Word files go through
# python-docx, PDFs through PyMuPDF, and pictures through Tesseract. Images never become the profile
# or background picture; that stays a manual upload.
import io
import os
import shutil
import subprocess
import tempfile
import zipfile

MAX_BYTES = 10 * 1024 * 1024        # upload size, checked before any parsing
MAX_UNZIPPED = 100 * 1024 * 1024    # Word files are zips; refuse ones that unpack to far more
MAX_PAGES = 20                      # PDF pages read
MAX_IMAGES = 20                     # embedded Word images read
MAX_CHARS = 30_000                  # text handed to the model
PAGE_DPI = 300                      # enough for body text; higher is slow for little gain
OCR_TIMEOUT = 90                    # seconds per image
TEXT_EXTS = (".txt", ".md")
IMAGE_EXTS = (".png", ".jpg", ".jpeg", ".webp")
ALLOWED = TEXT_EXTS + (".docx", ".pdf") + IMAGE_EXTS
# The first bytes of each type, so a renamed file is refused before a parser sees it.
SIGNATURES = {".docx": b"PK\x03\x04", ".pdf": b"%PDF", ".png": b"\x89PNG", ".jpg": b"\xff\xd8\xff",
              ".jpeg": b"\xff\xd8\xff", ".webp": b"RIFF"}
WINDOWS_DEFAULT = r"C:\Program Files\Tesseract-OCR\tesseract.exe"


class OCRError(ValueError):
    pass


def tesseract_path():
    # MISO_TESSERACT wins, then PATH, then the default Windows installer location.
    env = os.environ.get("MISO_TESSERACT")
    if env and os.path.isfile(env):
        return env
    found = shutil.which("tesseract")
    if found:
        return found
    if os.name == "nt" and os.path.isfile(WINDOWS_DEFAULT):
        return WINDOWS_DEFAULT
    return None


def _ocr_bytes(data, suffix):
    exe = tesseract_path()
    if exe is None:
        raise OCRError("OCR is not installed on this server. Install Tesseract, then try again (see the README).")
    fd, path = tempfile.mkstemp(suffix=suffix)
    try:
        with os.fdopen(fd, "wb") as f:
            f.write(data)
        try:
            run = subprocess.run([exe, path, "stdout", "-l", "eng", "--psm", "4"],
                                 capture_output=True, timeout=OCR_TIMEOUT, check=False)
        except subprocess.TimeoutExpired:
            raise OCRError("OCR took too long on one image.")
        if run.returncode != 0:
            raise OCRError("OCR failed on one image.")
        return run.stdout.decode("utf-8", errors="replace").strip()
    finally:
        os.remove(path)


def _ocr_or_note(data, suffix, notes):
    # For files that mix text and pictures: one missing OCR should not sink the whole file.
    try:
        return _ocr_bytes(data, suffix)
    except OCRError as err:
        if str(err) not in notes:
            notes.append(str(err))
        return ""


def _read_text(data):
    if data[:2] in (b"\xff\xfe", b"\xfe\xff"):
        return data.decode("utf-16", errors="replace")
    return data.decode("utf-8-sig", errors="replace")


def _read_docx(data, notes):
    import docx
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            if sum(i.file_size for i in z.infolist()) > MAX_UNZIPPED:
                raise OCRError("That Word file is too large once unpacked.")
    except zipfile.BadZipFile:
        raise OCRError("That Word file could not be read.")
    try:
        doc = docx.Document(io.BytesIO(data))
    except Exception:
        raise OCRError("That Word file could not be read.")
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    blobs = [rel.target_part.blob for rel in doc.part.rels.values() if rel.reltype.endswith("/image")]
    for blob in blobs[:MAX_IMAGES]:
        parts.append(_ocr_or_note(blob, ".png", notes))
    if len(blobs) > MAX_IMAGES:
        notes.append(f"Only the first {MAX_IMAGES} images were read.")
    return "\n\n".join(p for p in parts if p and p.strip())


def _read_pdf(data, notes):
    import pymupdf
    try:
        doc = pymupdf.open(stream=data, filetype="pdf")
    except Exception:
        raise OCRError("That PDF could not be read.")
    try:
        if doc.needs_pass:
            raise OCRError("That PDF is password protected.")
        parts, scanned = [], 0
        for index, page in enumerate(doc):
            if index >= MAX_PAGES:
                notes.append(f"Only the first {MAX_PAGES} pages were read.")
                break
            text = page.get_text().strip()
            if len(text) >= 20:
                parts.append(text)
                continue
            # No text layer, so the page is a picture of text: render it and OCR it.
            # Guard against decompression bombs: skip pages that would render to an oversized pixmap.
            rect = page.rect
            px_w = int(rect.width * PAGE_DPI / 72)
            px_h = int(rect.height * PAGE_DPI / 72)
            if px_w * px_h > 25_000_000:
                notes.append(f"Page {index + 1} is too large to render for OCR and was skipped.")
                continue
            scanned += 1
            parts.append(_ocr_or_note(page.get_pixmap(dpi=PAGE_DPI).tobytes("png"), ".png", notes))
        if scanned:
            notes.append(f"{scanned} page(s) had no text layer and were read with OCR.")
    finally:
        doc.close()
    return "\n\n".join(p for p in parts if p and p.strip())


def extract(filename, data):
    # Returns (text, notes). Raises OCRError with a message the admin can act on.
    ext = os.path.splitext((filename or "").lower())[1]
    if ext not in ALLOWED:
        raise OCRError("Use a .txt, .md, .docx, .pdf, .png, .jpg or .webp file.")
    if not data:
        raise OCRError("The file is empty.")
    if len(data) > MAX_BYTES:
        raise OCRError("Files can be up to 10 MB.")
    if ext in SIGNATURES and not data.startswith(SIGNATURES[ext]):
        raise OCRError("The file contents do not match its extension.")
    notes = []
    if ext in TEXT_EXTS:
        text = _read_text(data)
    elif ext == ".docx":
        text = _read_docx(data, notes)
    elif ext == ".pdf":
        text = _read_pdf(data, notes)
    else:
        text = _ocr_bytes(data, ext)
    text = text.strip()
    if not text:
        raise OCRError("No text was found in that file.")
    if len(text) > MAX_CHARS:
        text = text[:MAX_CHARS]
        notes.append(f"The text was cut to {MAX_CHARS:,} characters.")
    return text, notes
